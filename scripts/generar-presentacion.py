# -*- coding: utf-8 -*-
"""
Genera la presentación del segundo corte.

    python scripts/generar-presentacion.py

Diseño pensado para videobeam: fondo claro, texto grande, alto contraste,
sin adornos. Cada diapositiva trae notas del orador y, donde corresponde,
el nombre exacto de la actividad de la aplicación a la que se pasa.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "materiales" / "Presentacion_Segundo_Corte.pptx"

TINTA = RGBColor(0x14, 0x18, 0x1F)
TINTA_SUAVE = RGBColor(0x4A, 0x54, 0x64)
ACENTO = RGBColor(0x1A, 0x4F, 0xA0)
PAPEL = RGBColor(0xFF, 0xFF, 0xFF)
BANDA = RGBColor(0xE8, 0xF0, 0xFC)
VERDE = RGBColor(0x1A, 0x6B, 0x3C)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANCO = prs.slide_layouts[6]


def fondo(slide, color=PAPEL):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def caja(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def parrafo(tf, texto, tam=22, negrita=False, color=TINTA, espacio=10, primero=False, vinieta=None):
    p = tf.paragraphs[0] if primero else tf.add_paragraph()
    p.space_after = Pt(espacio)
    r = p.add_run()
    r.text = (f"{vinieta}  {texto}" if vinieta else texto)
    r.font.size = Pt(tam)
    r.font.bold = negrita
    r.font.color.rgb = color
    r.font.name = "Calibri"
    return p


def etiqueta(slide, texto):
    tf = caja(slide, 0.8, 0.45, 11.7, 0.5)
    parrafo(tf, texto.upper(), tam=14, negrita=True, color=ACENTO, primero=True)


def titulo(slide, texto, y=0.95):
    tf = caja(slide, 0.8, y, 11.7, 1.3)
    parrafo(tf, texto, tam=38, negrita=True, primero=True)


def notas(slide, texto):
    slide.notes_slide.notes_text_frame.text = texto


def diapositiva(seccion, tit, puntos, nota, cierre=None, tipo="normal"):
    s = prs.slides.add_slide(BLANCO)
    fondo(s, BANDA if tipo == "transicion" else PAPEL)
    etiqueta(s, seccion)
    titulo(s, tit)
    tf = caja(s, 0.9, 2.45, 11.5, 4.2)
    for i, p in enumerate(puntos):
        if isinstance(p, tuple):
            texto, tam, neg = p
            parrafo(tf, texto, tam=tam, negrita=neg, primero=(i == 0), espacio=14)
        else:
            parrafo(tf, p, tam=22, color=TINTA_SUAVE, primero=(i == 0), espacio=14, vinieta="—")
    if cierre:
        tfc = caja(s, 0.9, 6.35, 11.5, 0.8)
        parrafo(tfc, cierre, tam=18, negrita=True, color=ACENTO if tipo != "transicion" else VERDE, primero=True)
    notas(s, nota)
    return s


def tabla(seccion, tit, encabezados, filas, nota, cierre=None):
    s = prs.slides.add_slide(BLANCO)
    fondo(s)
    etiqueta(s, seccion)
    titulo(s, tit)
    n_f, n_c = len(filas) + 1, len(encabezados)
    alto = min(4.0, 0.5 + 0.48 * n_f)
    forma = s.shapes.add_table(n_f, n_c, Inches(0.9), Inches(2.45), Inches(11.5), Inches(alto))
    t = forma.table
    def celda(fila, col, texto, tam, negrita, color):
        c = t.cell(fila, col)
        c.text = str(texto) or " "
        pr = c.text_frame.paragraphs[0]
        if not pr.runs:          # una celda vacía no crea run: hay que forzarlo
            pr.add_run().text = " "
        pr.runs[0].font.size = Pt(tam)
        pr.runs[0].font.bold = negrita
        pr.runs[0].font.color.rgb = color

    for j, h in enumerate(encabezados):
        celda(0, j, h, 16, True, TINTA)
    for i, fila in enumerate(filas, start=1):
        for j, v in enumerate(fila):
            celda(i, j, v, 14, False, TINTA_SUAVE)
    if cierre:
        tfc = caja(s, 0.9, 2.55 + alto, 11.5, 1.0)
        parrafo(tfc, cierre, tam=18, negrita=True, color=ACENTO, primero=True)
    notas(s, nota)
    return s


# ============================== 1. Portada ==================================

s = prs.slides.add_slide(BLANCO)
fondo(s, BANDA)
tf = caja(s, 1.0, 2.0, 11.3, 3.2)
parrafo(tf, "INTRODUCCIÓN A LA INGENIERÍA INFORMÁTICA", tam=18, negrita=True, color=ACENTO, primero=True)
parrafo(tf, "Segundo corte · Sesiones 7 a 11", tam=46, negrita=True, espacio=16)
parrafo(tf, "Del proceso de desarrollo a los paradigmas de programación", tam=26, color=TINTA_SUAVE)
tf2 = caja(s, 1.0, 5.6, 11.3, 1.2)
parrafo(tf2, "Institución Universitaria Colegio Mayor del Cauca · Popayán", tam=17, color=TINTA_SUAVE, primero=True)
parrafo(tf2, "Ing. Ángela Patricia Paz Guañarita", tam=17, color=TINTA_SUAVE, espacio=0)
notas(s, "Abrir la sesión ubicando el corte dentro del semestre: vienen del primer parcial "
         "(ciclo de vida y tendencias) y ahora van a trabajar el proceso, la calidad y los "
         "paradigmas. Anunciar que todo el corte tiene una aplicación de apoyo donde la teoría "
         "de cada sesión viene con actividades que se corrigen solas.")

diapositiva(
    "Cómo vamos a trabajar",
    "Teoría proyectada, práctica en los equipos",
    [
        ("Seis estaciones en la aplicación del curso, una por bloque de contenido.", 24, False),
        ("Cada estación: qué aprenderás · comprende · mira un ejemplo · practica jugando · ponlo a prueba · qué te llevas.", 22, False),
        ("57 actividades que se verifican solas, con pistas progresivas y explicación de por qué falla cada opción.", 22, False),
        ("Un laboratorio de pseudocódigo que ejecuta el código de verdad, con la traza de las variables.", 22, False),
        ("El avance se guarda en la cuenta de cada quien: se puede empezar en la sala y seguir desde el celular.", 22, False),
    ],
    "Explicar la mecánica de las cinco sesiones: 25 minutos de teoría proyectada, 35 de práctica "
    "en la aplicación, 20 de puesta en común sobre lo que más falló. Pedirles que creen la cuenta "
    "hoy mismo: nombre, apellido, grupo, correo y contraseña. El código estudiantil es opcional y "
    "no sirve para entrar.",
    cierre="El ritmo de cada sesión: 25 min de teoría · 35 min de práctica · 20 min de puesta en común",
)

# ============================== SESIÓN 7 ====================================

diapositiva(
    "Sesión 7 · El proceso de desarrollo de software",
    "Qué vamos a poder hacer al terminar",
    [
        "Separar el problema de la solución en un caso real del campus.",
        "Escribir un requisito que otra persona pueda comprobar sin preguntarnos nada.",
        "Reconocer qué aporta cada actividad del proceso y qué rol la ejerce.",
        "Elegir el alcance de una primera entrega cuando la capacidad no alcanza para todo.",
        "Decidir qué hacer cuando llega un cambio a mitad del proyecto.",
    ],
    "Escribir los objetivos en el tablero y dejarlos visibles toda la sesión. Volver a ellos en la "
    "puesta en común: preguntar cuáles sienten que ya cumplen.",
)

diapositiva(
    "Sesión 7",
    "Producto y proceso no son lo mismo",
    [
        ("PRODUCTO: lo que se entrega. Ejecutable, API, app, documentación, datos.", 24, True),
        ("PROCESO: la secuencia de actividades que lo produce.", 24, True),
        "Dos equipos pueden entregar el mismo producto con procesos muy distintos.",
        "Esa diferencia explica cuánto cuesta cambiarlo, cuántos defectos llegan al usuario y si el equipo puede mantenerlo cuando quien lo escribió ya no está.",
        "Un proceso transforma una necesidad en un sistema Y en la evidencia de que lo resuelve.",
    ],
    "Insistir en la palabra evidencia: es lo que separa a un ingeniero de alguien que programa bien. "
    "Preguntarles: si entregan un trabajo que funciona pero no pueden mostrar cómo lo comprobaron, "
    "¿qué están entregando? Una promesa.",
    cierre="Una funcionalidad sin pruebas y sin requisitos escritos no es un entregable: es una promesa.",
)

tabla(
    "Sesión 7",
    "Las cinco actividades marco",
    ["Actividad", "Pregunta que responde", "Salida típica"],
    [
        ["Comunicación", "¿Qué necesita quién y para qué?", "Requisitos, historias, glosario"],
        ["Planeación", "¿Cómo, con quién, cuándo, con qué riesgos?", "Cronograma, estimaciones, riesgos"],
        ["Modelado", "¿Cómo se estructura antes de codificar?", "Diagramas, modelo de datos"],
        ["Construcción", "¿El código hace lo acordado?", "Código, pruebas, build"],
        ["Despliegue", "¿El usuario lo tiene y qué opina?", "Versión liberada, manual, feedback"],
    ],
    "Encima corren las actividades sombrilla, que nunca terminan: gestión de riesgos, aseguramiento "
    "de calidad, control de versiones, medición y revisiones. Decirles la frase clave: cuando un "
    "proyecto se cae, casi siempre falló una sombrilla, no la codificación.",
    cierre="Encima de todas: riesgos, calidad, control de versiones, medición y revisiones.",
)

diapositiva(
    "Sesión 7",
    "Los roles son decisiones, no cargos",
    [
        ("Analista de requisitos — traduce «que sea rápido» a un umbral medible.", 21, False),
        ("Arquitecto — decide la estructura. Sus decisiones son las más caras de revertir.", 21, False),
        ("Desarrollador — implementa y prueba su unidad.", 21, False),
        ("Ingeniero de pruebas (QA) — diseña experimentos que puedan fallar.", 21, False),
        ("DevOps — convierte «funciona en mi máquina» en «funciona en producción».", 21, False),
        ("Product Owner — decide qué NO se hace en esta versión.", 21, False),
        ("Diseñador UX — valida con usuarios reales, con evidencia, no con gustos.", 21, False),
    ],
    "En su proyecto de curso, de tres integrantes, los siete roles siguen existiendo: alguien "
    "tendrá que priorizar y alguien tendrá que escribir los criterios de aceptación. Si nadie lo "
    "hace explícitamente, esas decisiones igual se toman, pero por omisión.",
    cierre="En un equipo de tres, los siete roles siguen existiendo.",
)

tabla(
    "Sesión 7",
    "Modelos de proceso: cuál, cuándo y por qué",
    ["Modelo", "Funciona cuando", "Se rompe cuando"],
    [
        ["Cascada", "Requisitos estables, regulados o licitados", "El cliente descubre lo que quería al ver el producto"],
        ["Incremental", "Se necesita valor temprano y el alcance se parte", "Las partes tienen dependencias fuertes"],
        ["Prototipos", "La interacción es el riesgo principal", "Exigen que el prototipo sea el producto"],
        ["Espiral", "Proyecto grande, largo, de alto riesgo", "No hay presupuesto real para iterar"],
        ["Ágil (Scrum)", "Los requisitos cambian y hay acceso al usuario", "Se adopta la ceremonia sin las prácticas"],
    ],
    "No existe un modelo superior: existe uno adecuado a un nivel de incertidumbre y a un costo de "
    "error. Los dos criterios son independientes: hay proyectos ágiles con verificación muy "
    "estricta. Desmontar aquí el mito de que en Scrum no se documenta.",
    cierre="A mayor incertidumbre, ciclo de retroalimentación más corto. A mayor costo de falla, verificación más formal.",
)

diapositiva(
    "Sesión 7",
    "Un requisito sirve si se puede hacer fallar",
    [
        ("Vago:  «El sistema debe ser rápido e intuitivo.»", 24, True),
        ("Comprobable:  «El 95 % de las consultas responde en menos de 2 s con 200 usuarios concurrentes.»", 24, True),
        "Un requisito verificable indica QUÉ se mide, EN QUÉ CONDICIONES y CUÁL es el umbral.",
        "Prueba de fuego: si no puedes imaginar un experimento cuyo resultado obligue a decir «esto no se cumplió», todavía no es un requisito.",
        "Restricción ≠ requisito: «debe correr en los equipos de la sala 2» limita la solución, no describe una función.",
    ],
    "Pedir tres ejemplos al grupo en voz alta y convertirlos entre todos. Es el punto que más se "
    "evalúa en el parcial y el que más se falla.",
    cierre="Magnitud + condición + umbral.",
)

diapositiva(
    "Sesión 7 · a la práctica",
    "Abran la estación 1: Cómo se construye el software",
    [
        ("Ordena el desarrollo — organizar las acciones del sistema de reserva de salas.", 21, False),
        ("¿Quién se encarga? — emparejar responsabilidades con roles.", 21, False),
        ("Requisito o frase vaga — clasificar y después arreglar una frase vaga.", 21, False),
        ("Construye una entrega — elegir el alcance con 8 puntos de capacidad.", 21, False),
        ("Llegó un cambio — decidir y ver la consecuencia de la decisión.", 21, False),
        ("Ponlo a prueba: Del problema al primer incremento (4 pasos).", 21, True),
    ],
    "35 minutos. Pasar por los puestos. Los que terminen antes tienen dos desafíos opcionales: "
    "«Mitos sobre los modelos de proceso» y «Proyectar con velocidad». En la puesta en común, "
    "proyectar «Ordena el desarrollo» y resolverlo entre todos discutiendo las dependencias.",
    cierre="Taller de cierre: un problema del campus, sus usuarios y tres requisitos comprobables.",
    tipo="transicion",
)

# ============================== SESIÓN 8 ====================================

diapositiva(
    "Sesión 8 · Calidad, funcionalidad y usabilidad",
    "Qué vamos a poder hacer al terminar",
    [
        "Ubicar ISO/IEC 25010 dentro de la familia SQuaRE y nombrar sus ocho características.",
        "Distinguir una función ausente, una función incorrecta y una difícil de usar.",
        "Calcular tasa de éxito y tiempo en tarea con datos de una prueba con usuarios.",
        "Convertir un requisito en una comprobación que otro podría ejecutar.",
    ],
    "Arrancar con la frase de apertura: «funciona» es una opinión hasta que alguien define contra "
    "qué se compara. Preguntar cómo sabrían si una app es buena, y anotar en el tablero cuántas "
    "respuestas son medibles.",
)

tabla(
    "Sesión 8",
    "La familia ISO/IEC 25000 (SQuaRE)",
    ["División", "Numeración", "Qué contiene"],
    [
        ["Gestión de calidad", "2500n", "Vocabulario y guía general"],
        ["Modelo de calidad", "2501n", "ISO/IEC 25010 (producto) y 25012 (datos)"],
        ["Medición", "2502n", "Métricas de calidad"],
        ["Requisitos", "2503n", "Especificación de requisitos de calidad"],
        ["Evaluación", "2504n", "ISO/IEC 25040: el proceso de evaluación"],
    ],
    "Pregunta típica de parcial: ¿cuál es la norma del modelo de calidad del producto? Respuesta: "
    "25010. Advertir la confusión más común: ISO 9001 es gestión de la calidad de una ORGANIZACIÓN; "
    "una empresa certificada en 9001 puede entregar software malísimo sin contradicción.",
    cierre="El modelo de características es ISO/IEC 25010. No es 25000 ni 25001 ni 9001.",
)

diapositiva(
    "Sesión 8",
    "Las ocho características del producto (edición 2011)",
    [
        ("1. Adecuación funcional     5. Fiabilidad", 24, True),
        ("2. Eficiencia de desempeño  6. Seguridad", 24, True),
        ("3. Compatibilidad           7. Mantenibilidad", 24, True),
        ("4. Usabilidad               8. Portabilidad", 24, True),
        "La revisión de 2023 reorganiza el modelo en nueve: renombra usabilidad como capacidad de interacción, separa flexibilidad e incorpora seguridad física (safety).",
        "En el curso trabajamos la edición de ocho características. El error que sí se penaliza es mezclar las dos listas.",
    ],
    "Ser explícita con esto: hay material en internet con las nueve características de 2023 y "
    "material con las ocho de 2011. Las dos son correctas, pero no se mezclan. En el parcial se "
    "pregunta por la de ocho.",
    cierre="Trabajamos el modelo de 2011. Señalamos el cambio de 2023 donde corresponde.",
)

tabla(
    "Sesión 8",
    "Tres fallas que parecen la misma y no lo son",
    ["Lo que dice el usuario", "Qué pasó en realidad", "Característica"],
    [
        ["«No puedo cancelar mi reserva»", "La función no existe", "Adecuación funcional · completitud"],
        ["«Cancelé la mía y se borró la de otro»", "La función existe y hace algo distinto", "Adecuación funcional · corrección"],
        ["«Sí se puede, pero nadie encuentra cómo»", "Existe, es correcta, está escondida", "Usabilidad · operabilidad"],
    ],
    "Importa porque cada una la arregla alguien distinto: la primera es alcance (hay que "
    "construirla), la segunda es un defecto (corregir la lógica), la tercera es diseño de "
    "interacción (el código está bien y aun así el usuario no logra su objetivo).",
    cierre="Cada una la arregla alguien distinto. Por eso hay que nombrarlas bien.",
)

diapositiva(
    "Sesión 8",
    "Usabilidad: de la opinión a la evidencia",
    [
        ("Tasa de éxito = quienes completan la tarea ÷ total × 100", 24, True),
        ("Tiempo en tarea = MEDIANA de los tiempos, solo de quienes completaron", 24, True),
        ("Errores por tarea = total de errores ÷ participantes", 24, True),
        "Los tiempos de quienes NO completaron no entran: alguien que abandonó a los 30 segundos haría parecer eficiente un formulario que nadie logra llenar.",
        "Con 5 a 10 participantes se descubre qué arreglar, pero no se puede afirmar «el 80 % de nuestros usuarios lo logra».",
    ],
    "Contarles el ejemplo completo de la aplicación: 10 estudiantes, tarea «reserve la sala 203», "
    "70 % de éxito, mediana 2:30, y la causa concreta: un campo de fecha que exige AAAA-MM-DD sin "
    "decirlo. La conversación pasa de «la gente se queja» a «tres de cada diez no logran reservar».",
    cierre="La usabilidad no se opina: se mide con personas haciendo una tarea concreta.",
)

diapositiva(
    "Sesión 8 · a la práctica",
    "Abran la estación 2: ¿Cuándo un software es bueno?",
    [
        ("Detective de calidad — clasificar fallos por característica afectada.", 21, False),
        ("Encuentra la fricción — usar un formulario con defectos reales y nombrarlos.", 21, False),
        ("Mejora la interfaz — elegir cambios y ver cómo se transforma el formulario.", 21, False),
        ("Mide la usabilidad — calcular tasa de éxito, mediana y errores con datos reales.", 21, False),
        ("¿Qué probarías? — emparejar requisitos con comprobaciones observables.", 21, False),
        ("Ponlo a prueba: Diagnóstico de calidad de una app real (4 pasos).", 21, True),
    ],
    "«Encuentra la fricción» es la actividad para proyectar: intentar completar el formulario en "
    "vivo, delante de todos, y que el grupo vaya nombrando los problemas. Advertir que dos de los "
    "problemas listados NO ocurren: marcarlos resta, porque reportar defectos inexistentes también "
    "cuesta tiempo al equipo.",
    cierre="Taller de cierre: auditar una app conocida con la lista de chequeo y justificar dos mejoras.",
    tipo="transicion",
)

# ======================= LEY DE MOORE Y SOCIEDAD ============================

diapositiva(
    "Sesiones 8 y 9 · TIC en la sociedad",
    "Ley de Moore: qué dice exactamente",
    [
        ("El número de transistores en un circuito integrado de costo mínimo se duplica aproximadamente cada dos años.", 26, True),
        "Gordon Moore, 1965; él mismo la revisó en 1975 al periodo de dos años.",
        "Es una observación empírica sobre una industria. No es una ley física ni un derecho adquirido.",
        "Se ha venido desacelerando por límites físicos (fugas, disipación) y económicos (el costo de cada nueva planta).",
    ],
    "Ojo con la formulación: habla de DENSIDAD DE TRANSISTORES, no de velocidad. Ese matiz es todo "
    "el contenido de las dos diapositivas siguientes y es la pregunta más frecuente del parcial.",
    cierre="Describe transistores. No describe la velocidad de tu programa.",
)

tabla(
    "Sesiones 8 y 9",
    "Lineal y exponencial: dónde se separan",
    ["Año", "Lineal (+100 cada 2 años)", "Exponencial (×2 cada 2 años)"],
    [
        ["0", "100", "100"],
        ["2", "200", "200"],
        ["4", "300", "400"],
        ["6", "400", "800"],
        ["8", "500", "1.600"],
        ["12", "700", "6.400"],
    ],
    "A los 2 años son idénticos; a los 12, uno vale 700 y el otro 6.400. La diferencia no es «más "
    "rápido»: en el lineal la diferencia entre puntos consecutivos es constante; en el exponencial "
    "esa diferencia también crece. Fórmula: valor final = inicial × 2^(tiempo ÷ periodo).",
    cierre="Seis duplicaciones multiplican por 64, no por 12.",
)

diapositiva(
    "Sesiones 8 y 9",
    "Qué NO se puede concluir de la Ley de Moore",
    [
        ("«Mi programa correrá el doble de rápido en dos años.»", 24, True),
        "Desde ~2005 la frecuencia de reloj se estancó. Los transistores adicionales van a más núcleos, más caché y aceleradores.",
        "Un programa de un solo hilo usa un núcleo: no se acelera por sí solo.",
        "Ley de Amdahl: la fracción forzosamente secuencial acota la aceleración máxima. Si el 20 % no se puede repartir, ningún número de núcleos da más de 5×.",
        ("«Como el hardware mejora, no hace falta optimizar.» — Falso, y además ignora a quien usa un equipo de hace ocho años.", 22, True),
    ],
    "El patrón del error siempre es el mismo: tomar una observación sobre una magnitud "
    "(transistores) y sacar de ella una conclusión sobre otra (velocidad, garantías futuras, "
    "necesidad de optimizar). Es el mismo salto que van a oír toda la carrera con la inteligencia "
    "artificial.",
    cierre="Más transistores ≠ programas más rápidos.",
)

tabla(
    "Sesiones 8 y 9",
    "Marco legal colombiano: dos leyes que se confunden",
    ["", "Ley 1273 de 2009", "Ley 1581 de 2012"],
    [
        ["Qué es", "Reforma del Código Penal", "Ley estatutaria de datos personales"],
        ["Qué hace", "Crea el bien jurídico de protección de la información (arts. 269A–269J)", "Regula el tratamiento de datos personales"],
        ["A quién obliga", "A quien ATACA un sistema o unos datos", "A quien CUSTODIA datos"],
        ["Ejemplo", "Entrar con la contraseña de otro, aunque no dañe nada", "Guardar códigos y correos sin autorización ni finalidad"],
        ["Quién vigila", "Fiscalía y jueces penales", "Superintendencia de Industria y Comercio"],
    ],
    "Principios de la 1581: legalidad, finalidad, libertad, veracidad, transparencia, acceso y "
    "circulación restringida, seguridad y confidencialidad. El que más se incumple en proyectos "
    "académicos es finalidad: se piden datos «por si acaso». Recordarles que su proyecto final "
    "recoge datos de compañeros y que la ley les aplica desde ya.",
    cierre="La 1273 castiga a quien ataca. La 1581 obliga a quien custodia. Ustedes casi siempre serán lo segundo.",
)

diapositiva(
    "Sesiones 8 y 9 · a la práctica",
    "Abran la estación 4: Informática y sociedad",
    [
        ("Simulador de duplicación — mover valor inicial, periodo y tiempo, y leer el gráfico.", 21, False),
        ("Predice antes de mover — anticipar, ejecutar y comparar. Equivocarse es el punto.", 21, False),
        ("Mito o conclusión válida — siete afirmaciones sobre tecnología y crecimiento.", 21, False),
        ("Diseña para todos — 10 puntos de esfuerzo para ampliar quién puede usar el sistema.", 21, False),
        ("¿Necesitas ese dato? — decidir qué pide un formulario y con qué finalidad.", 21, False),
        ("Ponlo a prueba: Un servicio digital para todo el campus (4 pasos).", 21, True),
    ],
    "Proyectar «Predice antes de mover» y pedir la predicción a mano alzada ANTES de ejecutar: el "
    "de la hoja de papel doblada 20 veces casi nadie lo acierta (son 105 metros). Ese momento de "
    "error compartido es el que hace entender el exponencial.",
    cierre="Taller de cierre: analizar un servicio digital del campus — beneficios, barreras y datos.",
    tipo="transicion",
)

# ============================== SESIÓN 9 ====================================

diapositiva(
    "Sesión 9 · Portabilidad, compatibilidad y robustez",
    "Qué vamos a poder hacer al terminar",
    [
        "Escribir las reglas de validación de un campo: tipo, rango y formato.",
        "Elegir los valores de frontera que hay que probar en cualquier restricción numérica.",
        "Separar compatibilidad, portabilidad y disponibilidad con casos concretos.",
        "Decidir cómo debe comportarse una aplicación cuando se cae la conexión.",
        "Redactar un reporte de defecto que otra persona pueda reproducir.",
    ],
    "Frase de apertura: el software no falla cuando el usuario hace lo esperado; falla en los "
    "bordes. Preguntarles qué es lo más raro que han escrito en un formulario.",
)

diapositiva(
    "Sesión 9",
    "Las tres clases de entrada que hay que probar",
    [
        ("NORMAL — valor típico, lejos de los bordes.  Ej.: 3", 24, True),
        ("FRONTERA — justo en el borde y justo afuera.  Ej.: 0, 1, 5, 6", 24, True),
        ("INVÁLIDA — tipo, formato o ausencia.  Ej.: «tres», 2.5, vacío", 24, True),
        "Regla que no falla: para «entre A y B» siempre hay cuatro casos obligatorios — A−1, A, B y B+1.",
        "Con esos cuatro se detecta cualquier error de operador de comparación (< en vez de <=).",
    ],
    "Ejemplo del campo «cantidad de reservas, entre 1 y 5»: si el código escribió n < 5 en vez de "
    "n <= 5, la prueba con 5 falla y las otras tres pasan. Por eso las fronteras se prueban en "
    "pareja: el último válido y el primero inválido.",
    cierre="Un solo campo numérico produce ocho casos de prueba. No es exageración: es el mínimo.",
)

tabla(
    "Sesión 9",
    "Compatibilidad y portabilidad: la pregunta que las separa",
    ["¿El sistema…", "Entonces es…", "Subcaracterísticas"],
    [
        ["…intercambia información con otro producto, o convive con él?", "Compatibilidad", "Interoperabilidad, coexistencia"],
        ["…se instala y funciona en otro entorno?", "Portabilidad", "Adaptabilidad, instalación, reemplazo"],
    ],
    "Caso que confunde a todo el mundo: «el CSV abre bien en Excel y en LibreOffice» es "
    "COMPATIBILIDAD (hay intercambio de información), aunque mencione dos programas. «La misma app "
    "corre en Windows y en Ubuntu» es PORTABILIDAD: no hay intercambio, hay mudanza. Proyectar los "
    "dos casos juntos y pedir que voten antes de dar la respuesta.",
    cierre="Intercambio o convivencia → compatibilidad. Mudanza de entorno → portabilidad.",
)

tabla(
    "Sesión 9",
    "Qué significa un nueve más de disponibilidad",
    ["Disponibilidad", "Caída máxima al mes", "Caída máxima al año"],
    [
        ["99 %", "7 h 18 min", "3 días 15 h"],
        ["99,5 %", "3 h 39 min", "1 día 19 h"],
        ["99,9 %", "43 min 50 s", "8 h 46 min"],
        ["99,99 %", "4 min 23 s", "52 min 36 s"],
    ],
    "Dos consecuencias: cada nueve adicional cuesta unas diez veces más (exige redundancia, "
    "despliegue sin interrupción y monitoreo continuo); y un compromiso sin ventana declarada no "
    "significa nada, porque «99,9 % en horario de biblioteca» y «99,9 % 24×7» son cosas muy "
    "distintas.",
    cierre="De 99,5 % a 99,9 % el margen de caída se reduce a la quinta parte.",
)

diapositiva(
    "Sesión 9",
    "Y si se cae la conexión a mitad",
    [
        ("(a) Pantalla en blanco y datos perdidos.  Inaceptable.", 23, True),
        ("(b) «Error de red» y datos perdidos.  Honesto pero cruel.", 23, True),
        ("(c) Explica qué pasó, conserva lo escrito y deja el envío en cola.  Correcto.", 23, True),
        "Regla que no se negocia: nunca se confirma «Guardado» si la escritura no se completó.",
        "Tres estados, no dos: guardando · guardado · pendiente de sincronización.",
    ],
    "La aplicación del curso aplica esta regla con sus propios intentos: si se cae la red, dice "
    "«Pendiente de sincronización» y reintenta con un identificador que impide contar dos veces. "
    "Vale la pena mostrarlo en vivo desconectando el wifi un momento.",
    cierre="Prometer éxito por adelantado produce el conflicto que el sistema existía para evitar.",
)

diapositiva(
    "Sesión 9 · a la práctica",
    "Abran la estación 3: Que no se rompa",
    [
        ("Rompe el formulario — hay cuatro defectos reales; hay que provocarlos, no adivinarlos.", 21, False),
        ("Explora los límites — elegir el conjunto mínimo y suficiente de valores.", 21, False),
        ("¿Compatible o portable? — clasificar ocho situaciones, dos de ellas tramposas.", 21, False),
        ("Se cayó la conexión — elegir cómo responde la aplicación y ver la consecuencia.", 21, False),
        ("Diseña las pruebas — completar una tabla de entrada, resultado esperado y motivo.", 21, False),
        ("Ponlo a prueba: Poner a prueba una función nueva (4 pasos).", 21, True),
    ],
    "«Rompe el formulario» es la que hay que proyectar: el validador se ejecuta de verdad y muestra "
    "en una tabla qué exigían las reglas frente a qué hizo el sistema. Un defecto solo se descubre "
    "cuando el estudiante envía una entrada que los hace discrepar. Pedir que griten entradas raras.",
    cierre="Taller de cierre: un reporte de defecto real que otro grupo pueda reproducir.",
    tipo="transicion",
)

# ============================== SESIÓN 10 ===================================

diapositiva(
    "Sesión 10 · Paradigmas de programación",
    "Qué vamos a poder hacer al terminar",
    [
        "Trazar la ejecución de un algoritmo paso a paso, sin ejecutarlo.",
        "Usar secuencia, condición y repetición para resolver un problema.",
        "Distinguir contador, acumulador, bandera y centinela, y saber cuándo usar cada uno.",
        "Escribir pseudocódigo que pase casos normales y casos límite.",
        "Explicar qué son clase, objeto, atributo y método con un ejemplo propio.",
    ],
    "Frase de apertura: un paradigma no es una sintaxis, es una forma de repartir la "
    "responsabilidad dentro de un programa. Esta sesión es la más práctica del corte: conviene "
    "acortar la teoría a 20 minutos y dar 40 de laboratorio.",
)

diapositiva(
    "Sesión 10",
    "Antes de escribir nada: entrada, proceso y salida",
    [
        ("¿Qué datos recibe?  ¿Qué hace con ellos?  ¿Qué produce?", 26, True),
        "Si no puedes responder las tres, el problema está mal planteado y cualquier código será un intento a ciegas.",
        "La salida se especifica con su FORMATO EXACTO: «3.5» y «El promedio es 3.5» no son lo mismo.",
        "Si dos personas leen tu enunciado y escriben salidas distintas, falta precisión.",
    ],
    "En los ejercicios de la aplicación la salida siempre viene con su formato exacto, y no es "
    "capricho: es para que el estudiante aprenda que un enunciado ambiguo produce implementaciones "
    "distintas y ninguna está equivocada.",
    cierre="Un algoritmo: finito, ordenado, no ambiguo, y termina.",
)

tabla(
    "Sesión 10",
    "Los cuatro patrones que resuelven casi todo",
    ["Patrón", "Responde", "Inicialización", "Dentro del ciclo"],
    [
        ["Contador", "¿Cuántos?", "c ← 0", "c ← c + 1 (bajo condición)"],
        ["Acumulador", "¿Cuánto en total?", "s ← 0", "s ← s + x"],
        ["Bandera", "¿Ocurrió alguna vez?", "b ← Falso", "b ← Verdadero cuando ocurre"],
        ["Centinela", "¿Hasta cuándo leo?", "Leer antes del ciclo", "Procesar y volver a leer"],
    ],
    "El error más frecuente NO es de sintaxis: es inicializar dentro del ciclo. Si c ← 0 queda "
    "adentro, se reinicia en cada vuelta y el resultado siempre será 0 o 1, sin que el programa "
    "falle nunca. Es un defecto silencioso, de los peores. Con el centinela el error opuesto: leer "
    "dentro del ciclo antes de comprobar, con lo cual el centinela se procesa como dato.",
    cierre="Contador crece de a 1. Acumulador crece según el valor leído. No son lo mismo.",
)

diapositiva(
    "Sesión 10",
    "La asignación no es una igualdad",
    [
        ("x ← x + 1", 34, True),
        "Se evalúa TODO lo que está a la derecha y el resultado se guarda en la variable de la izquierda.",
        "En matemáticas, x = x + 1 no tiene solución. En programación es la operación más común del mundo.",
        "Trazar a mano es la habilidad que permite encontrar un error sin ejecutar, y es lo que se pregunta en el parcial.",
    ],
    "Hacer en vivo la traza del intercambio de dos variables con una temporal (t ← a; a ← b; b ← t) "
    "y preguntar qué pasa si se quita la temporal. En la aplicación, la actividad «Sigue las "
    "variables» calcula la traza con el intérprete: no hay respuestas escritas a mano que se puedan "
    "desactualizar.",
    cierre="Tres estructuras de control: secuencia, selección y repetición. Con eso se construye todo.",
)

tabla(
    "Sesión 10",
    "De la programación estructurada a los objetos",
    ["Clase", "Atributos (sustantivos)", "Métodos (verbos)"],
    [
        ["Sala", "código, capacidad, tiene proyector", "reservar(franja), liberar(), estaDisponible()"],
        ["Reserva", "sala, estudiante, fecha, hora inicio, hora fin", "cancelar(), duracion(), seSolapaCon(otra)"],
        ["Estudiante", "nombres, código, programa", "reservasActivas(), puedeReservar()"],
    ],
    "La prueba que separa clase de objeto: si puedes contarlos o señalar uno, es objeto. «Sala» no "
    "se puede contar, así que es clase; «la sala 203» sí, así que es objeto; y 30 no es ninguna de "
    "las dos, es el valor de un atributo. Error clásico del parcial: tratar «Sala203» como si fuera "
    "una clase distinta.",
    cierre="Lo que cambia entre paradigmas no es la lógica: es quién responde por ella.",
)

diapositiva(
    "Sesión 10 · a la práctica",
    "Abran la estación 5: Pensar como programador",
    [
        ("Sigue las variables — completar la traza paso a paso.", 21, False),
        ("Arma el algoritmo — ordenar bloques; el resultado SE EJECUTA contra los casos.", 21, False),
        ("Completa la decisión — elegir condiciones y probar las fronteras exactas.", 21, False),
        ("Repara el ciclo — encontrar dos defectos: uno ruidoso y uno silencioso.", 21, False),
        ("Contador o acumulador — completar y ver la traza de cada variable.", 21, False),
        ("Del objeto a la clase — agrupar objetos y separar atributos de métodos.", 21, False),
        ("Ponlo a prueba: Reto integrador · Informe de notas.", 21, True),
    ],
    "Insistir en que los juegos EJECUTAN el algoritmo: si su orden funciona, pasa, aunque no "
    "coincida con el que yo tenía en mente. Y si el orden «correcto» no produce la salida, no "
    "pasa. Para quienes terminen antes: el desafío opcional de lectura con centinela. El "
    "laboratorio queda abierto para el trabajo del tablero.",
    cierre="Taller de cierre: un algoritmo propio en pseudocódigo, con sus decisiones explicadas.",
    tipo="transicion",
)

# ============================== SESIÓN 11 ===================================

diapositiva(
    "Sesión 11",
    "Repaso e instrucciones del segundo parcial",
    [
        ("Primeros 45 minutos: repaso dirigido por lo que más falló.", 24, True),
        "Revisar en el panel docente las actividades con mayor porcentaje sin resolver y proyectar esas.",
        "Resolver en vivo las dos o tres que concentran los errores, pidiendo el razonamiento antes de la respuesta.",
        ("Últimos 15 minutos: instrucciones del parcial.", 24, True),
        "Alcance: sesiones 7 a 10. Formato: selección múltiple con justificación, un caso de calidad y un algoritmo en pseudocódigo.",
    ],
    "Traer impresa o proyectada la tabla del panel docente con las actividades más atascadas. "
    "Recordarles que el reto final de práctica (estación 6) es el ensayo general: mismo alcance y "
    "mismo tipo de preguntas, con cronómetro opcional de 45 minutos, y que al terminarlo reciben un "
    "informe por concepto con el enlace exacto a lo que conviene repasar.",
    cierre="El reto final de práctica no es el parcial ni genera nota: es el ensayo.",
)

diapositiva(
    "Sesión 11 · antes del parcial",
    "Cómo prepararse con la aplicación",
    [
        ("1. Hagan el reto final de práctica SIN cronómetro la primera vez.", 22, True),
        ("2. Lean el informe por concepto que sale al terminar.", 22, True),
        ("3. Vuelvan solo a las secciones que el informe señale. No relean todo.", 22, True),
        ("4. Repitan el reto CON cronómetro de 45 minutos.", 22, True),
        "En «Mi progreso» cada quien ve sus intentos, las pistas que usó y sus recomendaciones de repaso.",
        "La nota orientativa sobre 5,0 es una referencia del avance, no la nota del corte.",
    ],
    "Dejar claro el límite de la nota orientativa: sale de las actividades base de la aplicación y "
    "excluye desafíos opcionales y talleres. La nota académica la define el parcial más los "
    "entregables. Si alguien pregunta si la app «pone nota», la respuesta es no.",
    cierre="Fallar en el ensayo es barato. Fallar en el parcial no.",
)

s = prs.slides.add_slide(BLANCO)
fondo(s, BANDA)
tf = caja(s, 1.0, 2.6, 11.3, 2.6)
parrafo(tf, "Lo que se llevan del corte", tam=40, negrita=True, primero=True)
parrafo(tf, "Un requisito sirve si se puede hacer fallar. La calidad se afirma con evidencia. "
            "Los defectos viven en los bordes. Una observación no es una promesa. "
            "Y un algoritmo sin casos de prueba tiene el mismo problema que un requisito sin umbral: "
            "nadie puede decir si está bien.", tam=22, color=TINTA_SUAVE, espacio=14)
notas(s, "Cerrar el corte conectando las cinco sesiones: todas responden a la misma pregunta —"
         "¿cómo sé que esto está bien?— y todas la responden igual: con evidencia que otro pueda "
         "reproducir. Es el hilo que sigue en el resto de la carrera.")

SALIDA.parent.mkdir(parents=True, exist_ok=True)
prs.save(SALIDA)
print(f"Presentación generada: {SALIDA}")
print(f"  {len(prs.slides.__iter__.__self__._sldIdLst)} diapositivas")
