# -*- coding: utf-8 -*-
"""
Genera la guía docente del segundo corte.

    python scripts/generar-guia.py

La guía describe la aplicación tal como está construida: los nombres de las
actividades, las estaciones y las mecánicas que se mencionan existen y se
pueden abrir. Si se agrega o renombra una actividad, hay que actualizar aquí.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

RAIZ = Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "materiales" / "Guia_Docente_Segundo_Corte.docx"

doc = Document()

# Márgenes y tipografía base
for s in doc.sections:
    s.top_margin = Inches(0.9)
    s.bottom_margin = Inches(0.9)
    s.left_margin = Inches(1.0)
    s.right_margin = Inches(1.0)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15


def h1(t):
    p = doc.add_heading(t, level=1)
    p.runs[0].font.color.rgb = RGBColor(0x1A, 0x4F, 0xA0)
    return p


def h2(t):
    p = doc.add_heading(t, level=2)
    p.runs[0].font.color.rgb = RGBColor(0x14, 0x18, 0x1F)
    return p


def h3(t):
    p = doc.add_heading(t, level=3)
    p.runs[0].font.color.rgb = RGBColor(0x4A, 0x54, 0x64)
    return p


def p(t, negrita=False, cursiva=False):
    par = doc.add_paragraph()
    r = par.add_run(t)
    r.bold = negrita
    r.italic = cursiva
    return par


def vinetas(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


def numerada(items):
    for i in items:
        doc.add_paragraph(i, style="List Number")


def tabla(encabezados, filas, anchos=None):
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(encabezados):
        c = t.rows[0].cells[j]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for fila in filas:
        cells = t.add_row().cells
        for j, v in enumerate(fila):
            cells[j].text = ""
            run = cells[j].paragraphs[0].add_run(str(v))
            run.font.size = Pt(10)
    doc.add_paragraph()
    return t


# ============================== Portada =====================================

t = doc.add_heading("Guía docente", level=0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Introducción a la Ingeniería Informática · Segundo corte · Sesiones 7 a 11")
r.bold = True
r.font.size = Pt(14)
sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub2.add_run(
    "Institución Universitaria Colegio Mayor del Cauca · Popayán\n"
    "Ing. Ángela Patricia Paz Guañarita"
).font.size = Pt(11)

doc.add_paragraph()
p(
    "Esta guía acompaña la aplicación del segundo corte y la presentación de 35 diapositivas. "
    "Las actividades que se nombran aquí existen en la aplicación con ese nombre exacto: se "
    "pueden abrir y proyectar tal cual.",
    cursiva=True,
)

doc.add_page_break()

# ============================== 1. Panorama =================================

h1("1. Panorama del corte")

p(
    "El corte cubre cinco sesiones del microcurrículo (semanas 7 a 11) y está organizado en seis "
    "estaciones dentro de la aplicación. Cada estación tiene siempre la misma estructura, para que "
    "el estudiante sepa qué esperar: qué aprenderás, comprende, mira un ejemplo, practica jugando, "
    "ponlo a prueba y qué te llevas."
)

tabla(
    ["Sesión", "Tema del microcurrículo", "Estación", "Actividades"],
    [
        ["7", "El proceso de desarrollo de software", "1. Cómo se construye el software", "11"],
        ["8", "Calidad SQuaRE, corrección funcional, usabilidad", "2. ¿Cuándo un software es bueno?", "11"],
        ["9", "Portabilidad, compatibilidad, robustez", "3. Que no se rompa", "11"],
        ["8 y 9", "Ley de Moore, TIC en la sociedad", "4. Informática y sociedad", "10"],
        ["10", "Paradigmas: del algoritmo a la POO", "5. Pensar como programador", "8"],
        ["11", "Repaso y segundo parcial", "6. Reto final de práctica", "6"],
    ],
)

p(
    "Además hay un laboratorio de pseudocódigo con seis ejemplos cargables, un glosario de 40 "
    "términos y las vistas «Mi progreso» (estudiante) y «Panel docente» (usted)."
)

h2("Qué hace distinta a esta aplicación")

vinetas(
    [
        "La teoría alcanza para resolver las actividades. Nada importante está escondido detrás de una respuesta incorrecta.",
        "Las actividades de código EJECUTAN el algoritmo del estudiante contra casos de prueba, incluidos los límites. No se comparan con una solución guardada, así que las soluciones alternativas válidas también pasan.",
        "Cada opción incorrecta explica por qué falla. Las opciones incorrectas están construidas sobre errores comprensibles, no son absurdos para facilitar el acierto.",
        "Se distinguen tres estados: intentada, completada y dominio demostrado (al primer intento y sin pistas). Abrir una tarjeta o marcar una lista no produce ninguno.",
        "El puntaje lo calcula el servidor a partir del catálogo, no el navegador del estudiante.",
    ]
)

# ============================== 2. Planeación ===============================

doc.add_page_break()
h1("2. Planeación por sesión")

p(
    "Todas las sesiones son de dos horas (120 minutos) y usan la misma secuencia: explicación, "
    "demostración, práctica y cierre. La distribución que sigue es una propuesta: si el grupo va "
    "más lento, recorte la práctica antes que la puesta en común, porque es ahí donde se corrigen "
    "los errores de concepto."
)

tabla(
    ["Momento", "Minutos", "Qué se hace"],
    [
        ["Apertura", "10", "Objetivos en el tablero y una pregunta que active lo previo."],
        ["Explicación", "25", "Teoría proyectada con las diapositivas de la sesión."],
        ["Demostración", "15", "Resolver en vivo una actividad de la aplicación, pensando en voz alta."],
        ["Práctica", "45", "Los estudiantes trabajan la estación en los equipos. Usted pasa por los puestos."],
        ["Cierre", "25", "Puesta en común sobre lo que más falló, revisado en el panel docente."],
    ],
)

SESIONES = [
    {
        "n": "Sesión 7",
        "titulo": "El proceso de desarrollo de software",
        "objetivo": "Que el estudiante distinga producto de proceso, reconozca las actividades y "
        "roles del desarrollo, escriba un requisito comprobable y decida el alcance de una "
        "primera entrega.",
        "diapositivas": "3 a 9",
        "demostracion": "«Ordena el desarrollo» (estación 1). Resuélvala en vivo preguntando, "
        "paso por paso, qué depende de qué. Es la que mejor muestra que el proceso es una red "
        "de dependencias y no una fila india.",
        "practica": [
            "Ordena el desarrollo · ordenar por dependencias, admite varios órdenes válidos.",
            "¿Quién se encarga? · emparejar responsabilidades con roles.",
            "Requisito o frase vaga · clasificar y después convertir una frase vaga.",
            "Construye una entrega · elegir el alcance con 8 puntos de capacidad.",
            "Llegó un cambio · decidir y ver la consecuencia antes de confirmar.",
            "Ponlo a prueba: Del problema al primer incremento (4 pasos).",
        ],
        "opcionales": ["Mitos sobre los modelos de proceso (V/F)", "Proyectar con velocidad (cálculo)"],
        "cierre": "Proyectar «Requisito o frase vaga» y discutir los dos casos que más se fallan: "
        "«el sistema debe ser seguro» (vaga) frente a «debe cumplir la Ley 1581» (restricción).",
        "taller": "Un problema del campus: enunciado sin tecnología, usuarios y tres requisitos comprobables.",
    },
    {
        "n": "Sesión 8",
        "titulo": "Calidad del software, corrección funcional y usabilidad",
        "objetivo": "Que el estudiante ubique ISO/IEC 25010 en la familia SQuaRE, distinga función "
        "ausente / incorrecta / difícil de usar, y mida usabilidad con datos.",
        "diapositivas": "10 a 16 (calidad) y 17 a 21 (Ley de Moore y sociedad, si alcanza)",
        "demostracion": "«Encuentra la fricción» (estación 2). Intente completar el formulario en "
        "vivo, delante del grupo, y deje que ellos vayan nombrando los problemas. Es incómodo a "
        "propósito y funciona muy bien como demostración.",
        "practica": [
            "Detective de calidad · clasificar fallos por característica afectada.",
            "Encuentra la fricción · formulario con cinco defectos reales y dos falsos.",
            "Mejora la interfaz · elegir cambios y ver el formulario transformarse.",
            "Mide la usabilidad · tasa de éxito, mediana y errores con datos de 10 participantes.",
            "¿Qué probarías? · emparejar requisitos con comprobaciones observables.",
            "Ponlo a prueba: Diagnóstico de calidad de una app real (4 pasos).",
        ],
        "opcionales": ["Ubicar la norma (quiz)", "Cada subcaracterística en su lugar (clasificar)"],
        "cierre": "Revisar «Mide la usabilidad»: el punto que más se falla es incluir en el tiempo "
        "en tarea a quienes NO completaron. Pregunte por qué eso haría parecer eficiente un "
        "formulario que nadie logra llenar.",
        "taller": "Auditar una app conocida con la lista de chequeo y justificar dos mejoras.",
    },
    {
        "n": "Sesión 9",
        "titulo": "Portabilidad, compatibilidad, robustez y medición",
        "objetivo": "Que el estudiante diseñe las entradas que hacen fallar un sistema, distinga "
        "compatibilidad de portabilidad y redacte un reporte de defecto reproducible.",
        "diapositivas": "22 a 27",
        "demostracion": "«Rompe el formulario» (estación 3). El validador se ejecuta de verdad: "
        "pida entradas raras a viva voz y muéstrelas en la tabla que compara lo que exigían las "
        "reglas con lo que hizo el sistema. Un defecto solo se descubre provocándolo.",
        "practica": [
            "Rompe el formulario · cuatro defectos reales que hay que provocar, no adivinar.",
            "Explora los límites · elegir el conjunto mínimo y suficiente de valores.",
            "¿Compatible o portable? · ocho situaciones, dos de ellas diseñadas para confundir.",
            "Se cayó la conexión · elegir el comportamiento correcto ante una falla de red.",
            "Diseña las pruebas · completar entrada, resultado esperado y motivo.",
            "Ponlo a prueba: Poner a prueba una función nueva (4 pasos).",
        ],
        "opcionales": ["Cuánto es un nueve más (cálculo)", "Anatomía de un reporte de defecto (ordenar)"],
        "cierre": "Proyectar el par de casos tramposos de «¿Compatible o portable?»: el CSV que "
        "abre en Excel y en LibreOffice (compatibilidad) frente a la app que corre en Windows y "
        "Ubuntu (portabilidad). Pida que voten antes de revelar.",
        "taller": "Un reporte de defecto real que otro grupo pueda reproducir sin preguntar.",
    },
    {
        "n": "Sesión 10",
        "titulo": "Paradigmas de programación: del algoritmo a la orientación a objetos",
        "objetivo": "Que el estudiante trace un algoritmo paso a paso, use las tres estructuras de "
        "control con los cuatro patrones, y explique clase, objeto, atributo y método.",
        "diapositivas": "28 a 33",
        "demostracion": "«Sigue las variables» (estación 5) y el laboratorio. Trace en el tablero "
        "el intercambio de dos variables con una temporal y después muestre la traza que produce "
        "el intérprete. Es la sesión donde conviene acortar la teoría a 20 minutos.",
        "practica": [
            "Sigue las variables · completar la traza paso a paso.",
            "Arma el algoritmo · ordenar bloques; el resultado se ejecuta contra 5 casos.",
            "Completa la decisión · elegir condiciones y probar las fronteras exactas.",
            "Repara el ciclo · dos defectos: un ciclo infinito y una inicialización mal puesta.",
            "Contador o acumulador · completar y ver la traza de cada variable.",
            "Del objeto a la clase · agrupar objetos y separar atributos de métodos.",
            "Ponlo a prueba: Reto integrador · Informe de notas.",
        ],
        "opcionales": ["Lectura con centinela (pseudocódigo)"],
        "cierre": "Revisar «Repara el ciclo». El defecto que importa no es el ciclo infinito —que "
        "se nota— sino la inicialización dentro del ciclo, que no falla nunca y devuelve siempre "
        "un resultado equivocado. Los defectos silenciosos son los que llegan a producción.",
        "taller": "Un algoritmo propio en pseudocódigo, con sus decisiones y casos límite explicados.",
    },
    {
        "n": "Sesión 11",
        "titulo": "Repaso y segundo parcial",
        "objetivo": "Cerrar los vacíos que muestren los datos y dejar claras las condiciones del parcial.",
        "diapositivas": "34 y 35",
        "demostracion": "Abrir el panel docente y proyectar la tabla «Actividades que más se "
        "atascan», filtrada por el grupo. Resolver en vivo las dos o tres primeras.",
        "practica": [
            "Reto final de práctica (estación 6): seis pasos sobre un caso nuevo, con cronómetro opcional de 45 minutos.",
            "Al terminar, cada estudiante recibe un informe por concepto con el enlace exacto a lo que conviene repasar.",
        ],
        "opcionales": [],
        "cierre": "Instrucciones del parcial: alcance (sesiones 7 a 10), formato y duración. "
        "Recordar que el reto final no es el parcial ni genera nota académica.",
        "taller": "—",
    },
]

for s in SESIONES:
    doc.add_page_break()
    h2(f"{s['n']} · {s['titulo']}")
    p("Objetivo de la sesión", negrita=True)
    p(s["objetivo"])
    p(f"Diapositivas: {s['diapositivas']}")

    h3("Demostración (15 min)")
    p(s["demostracion"])

    h3("Práctica en la aplicación (45 min)")
    vinetas(s["practica"])
    if s["opcionales"]:
        p("Para quienes terminen antes (desafíos opcionales, no obligatorios):", negrita=True)
        vinetas(s["opcionales"])

    h3("Cierre (25 min)")
    p(s["cierre"])

    if s["taller"] != "—":
        h3("Taller de la sesión")
        p(s["taller"])
        p(
            "El taller no lo califica la aplicación: es un entregable de la plataforma del curso. "
            "La aplicación sí guarda el borrador para que el grupo no pierda el trabajo, y muestra "
            "la rúbrica desde el principio.",
            cursiva=True,
        )

# ============================== 3. Errores frecuentes =======================

doc.add_page_break()
h1("3. Errores frecuentes y cómo explicarlos")

p(
    "Estos son los errores que la aplicación está diseñada para provocar y corregir. Conviene "
    "anticiparlos: si aparecen en clase, ya tiene la explicación lista."
)

tabla(
    ["Error", "Por qué ocurre", "Cómo explicarlo"],
    [
        [
            "Confundir restricción con requisito",
            "Las dos se escriben como obligaciones del sistema.",
            "El requisito describe algo que el sistema HACE; la restricción LIMITA la solución sin describir ninguna función.",
        ],
        [
            "Creer que en ágil no se documenta",
            "Se convierte una preferencia del manifiesto («por encima de») en una prohibición.",
            "Si una opción del parcial prohíbe algo en ágil, sospechar de ella. Se documenta lo que se va a leer.",
        ],
        [
            "Confundir ISO 9001 con ISO/IEC 25010",
            "Las dos hablan de calidad.",
            "9001 certifica el sistema de gestión de una empresa; 25010 describe el producto. Una empresa certificada en 9001 puede entregar software malísimo sin contradicción.",
        ],
        [
            "Mezclar el modelo de 8 características (2011) con el de 9 (2023)",
            "Hay material de ambas ediciones circulando.",
            "Trabajamos la de 2011. Lo que se penaliza no es usar una u otra: es mezclarlas en la misma lista.",
        ],
        [
            "Llamar «usabilidad» a un defecto de corrección funcional",
            "El usuario reporta ambos como «no funciona».",
            "Preguntar: ¿la función existe? ¿da el resultado acordado? Si existe y es correcta pero el usuario no llega a ella, ahí sí es usabilidad.",
        ],
        [
            "Incluir en el tiempo en tarea a quienes no completaron",
            "Parece que todos los datos deben contar.",
            "Quien abandonó a los 30 segundos produce un tiempo bajísimo. Incluirlo hace parecer eficiente un formulario que nadie logra llenar.",
        ],
        [
            "Probar solo valores razonables",
            "Se prueba para confirmar que funciona, no para hacerlo fallar.",
            "Para «entre A y B» siempre hay cuatro casos: A−1, A, B y B+1. Ahí vive el error de < contra <=.",
        ],
        [
            "Confundir compatibilidad con portabilidad",
            "Las dos suenan a «funcionar en otro lado».",
            "¿Intercambia o convive con otro producto? Compatibilidad. ¿Se muda a otro entorno? Portabilidad.",
        ],
        [
            "Decir que la Ley de Moore promete programas más rápidos",
            "Hasta 2005 transistores y frecuencia avanzaron juntos.",
            "Describe densidad de transistores. Desde ~2005 se usan en más núcleos, y un programa de un hilo no se acelera solo.",
        ],
        [
            "Pedir datos «por si acaso»",
            "Costumbre de los formularios institucionales.",
            "Principio de finalidad: si no se puede escribir para qué se usará el dato, no se pide. El dato que no se recoge no se puede filtrar.",
        ],
        [
            "Creer que reemplazar el nombre por el código anonimiza",
            "El código parece impersonal.",
            "El código identifica a una persona determinada: eso es seudonimización, y sigue siendo dato personal.",
        ],
        [
            "Inicializar el contador dentro del ciclo",
            "Se escribe en el orden en que se piensa.",
            "Es un defecto silencioso: el programa no falla y siempre devuelve 0 o 1. Trazar dos vueltas a mano lo hace evidente.",
        ],
        [
            "Sumar el centinela",
            "Se lee al principio del cuerpo del ciclo.",
            "Lectura adelantada: leer una vez ANTES del Mientras y volver a leer al FINAL del cuerpo.",
        ],
        [
            "Tratar «Sala203» como una clase",
            "Cada caso concreto parece distinto.",
            "Si se puede contar o señalar, es objeto. Si hiciera falta una clase por sala, habría que reescribir el programa cada vez que abran un aula.",
        ],
        [
            "Confundir atributo con método por el nombre",
            "«tieneProyector» suena a pregunta; «duracion()» suena a dato.",
            "Atributo = guarda estado. Método = lo calcula o lo cambia. duracion() se deriva de las horas: guardarla crearía inconsistencias.",
        ],
    ],
)

# ============================== 4. Diferencias de nivel =====================

doc.add_page_break()
h1("4. Estudiantes con y sin experiencia previa")

p(
    "El grupo mezcla primíparos sin ninguna experiencia con estudiantes que vienen de programas "
    "de tecnología. La aplicación no los separa ni los etiqueta en público: todos recorren la "
    "misma ruta y los desafíos opcionales están a la vista de todos, marcados como opcionales."
)

h2("Apoyo para quienes no tienen experiencia")

vinetas(
    [
        "Las pistas son progresivas y cuestan puntaje, pero el piso es el 40 %: equivocarse varias veces sigue valiendo la pena. Dígalo en voz alta el primer día, porque muchos no piden ayuda por miedo a perder todo.",
        "En cada tarjeta de teoría hay un botón «Quiero entender mejor» que amplía el concepto sin salir de la página. Los términos subrayados abren su definición al tocarlos.",
        "Para el pseudocódigo: pídales que primero escriban entrada, proceso y salida en el cuaderno. La mayoría de los bloqueos vienen de no saber qué debe producir el algoritmo.",
        "En el laboratorio, cargar el ejemplo «Contador y acumulador juntos» y modificarlo es mucho más productivo que empezar de cero.",
        "«Ver una solución» existe en las actividades de pseudocódigo: vale 0 puntos pero deja estudiarla. Es preferible a que abandonen.",
    ]
)

h2("Extensiones para quienes avanzan rápido")

vinetas(
    [
        "Los desafíos opcionales de cada estación (10 en total). No cuentan para la nota orientativa y no bloquean la ruta principal.",
        "Pedirles que encuentren una segunda solución válida a «Arma el algoritmo» y la ejecuten: descubren que hay más de un orden correcto.",
        "Pedirles que rompan un algoritmo ya resuelto: cambiar >= por > y predecir qué caso de prueba fallará antes de ejecutarlo.",
        "Que escriban un caso de prueba adicional para el reto integrador y expliquen qué clase de entrada cubre.",
        "Que actúen como monitores en la práctica: explicar es la mejor forma de consolidar.",
    ]
)

# ============================== 5. Evaluación ===============================

doc.add_page_break()
h1("5. Propuesta de evaluación del corte")

p(
    "El microcurrículo asigna 35 % al segundo parcial. Esta es una propuesta de distribución que "
    "usa la aplicación como seguimiento, sin que el juego reemplace la evaluación."
)

tabla(
    ["Componente", "Peso", "Cómo se obtiene"],
    [
        ["Segundo parcial escrito", "60 %", "Examen de la sesión 11, sobre las sesiones 7 a 10."],
        ["Talleres de clase (4)", "25 %", "Entregables de las estaciones 1 a 5, con la rúbrica de cada uno."],
        ["Avance en la aplicación", "15 %", "Actividades base completadas. Se toma la nota orientativa que muestra el panel docente."],
    ],
)

p("Sobre el 15 % de la aplicación:", negrita=True)
vinetas(
    [
        "La nota orientativa se calcula como puntos obtenidos ÷ puntos posibles × 5,0, contando SOLO las actividades de nivel base (4.420 puntos posibles).",
        "Quedan excluidos los desafíos opcionales, los talleres y el laboratorio.",
        "El puntaje de cada intento lo recalcula el servidor a partir del catálogo: el navegador informa evidencia (si acertó, cuántas pistas abrió, cuántos intentos falló), no el puntaje.",
        "La corrección de cada actividad sí ocurre en el navegador. Para el 15 % de seguimiento es suficiente; no lo use como única evidencia de una nota alta sin contrastarla con el parcial.",
    ]
)

h2("Rúbrica común de los talleres")

tabla(
    ["Criterio", "Peso", "Qué se observa"],
    [
        ["Precisión conceptual", "40 %", "Usa los términos del corte con el significado correcto y los aplica al caso elegido."],
        ["Verificabilidad", "30 %", "Lo entregado se puede comprobar: otro grupo podría ejecutar la prueba o reproducir el defecto."],
        ["Justificación", "30 %", "Cada decisión dice por qué, y descarta explícitamente una alternativa."],
    ],
)

p(
    "Cada taller tiene además su propia rúbrica específica, visible para el estudiante dentro de "
    "la aplicación desde antes de empezar.",
    cursiva=True,
)

# ============================== 6. Banco de preguntas =======================

doc.add_page_break()
h1("6. Doce preguntas de parcial, con respuesta y justificación")

PREGUNTAS = [
    (
        "Sesión 7",
        "¿Cuál de los siguientes enunciados es un requisito verificable tal como está escrito?",
        [
            "El sistema debe ser rápido e intuitivo.",
            "El sistema debe usar las mejores prácticas de la industria.",
            "El 95 % de las consultas de saldo debe responder en menos de 800 ms con 200 usuarios concurrentes.",
            "El sistema debe ser moderno y escalable.",
        ],
        "c",
        "Un requisito verificable indica magnitud (800 ms), condición (200 concurrentes) y umbral "
        "(95 %). «Rápido», «intuitivo» y «moderno» no definen ningún experimento que pueda fallar, "
        "así que no se pueden cumplir ni incumplir de forma demostrable.",
    ),
    (
        "Sesión 7",
        "Un proyecto en cascada lleva cuatro meses cuando un decreto cambia la fórmula de "
        "liquidación que el sistema calcula. ¿Cuál es el diagnóstico correcto?",
        [
            "Fue un error usar cascada: con Scrum el decreto no habría afectado el proyecto.",
            "El modelo no evita el cambio; lo que falló fue no tratar la volatilidad normativa como riesgo y no aislar la fórmula como componente configurable.",
            "El cambio debe rechazarse porque el alcance estaba firmado.",
            "Es un defecto de calidad del proveedor y debe corregirse sin costo.",
        ],
        "b",
        "Ningún modelo de proceso impide que cambie la ley. Lo que absorbe el cambio es la gestión "
        "de riesgos (actividad sombrilla) y una decisión de arquitectura que aísle lo que se sabe "
        "volátil. Culpar al modelo es la respuesta fácil y la equivocada. Y no es un defecto: el "
        "sistema hace exactamente lo acordado; lo que cambió fue el acuerdo.",
    ),
    (
        "Sesión 8",
        "¿Qué norma de la familia SQuaRE define el modelo de calidad del producto de software?",
        ["ISO/IEC 25000", "ISO/IEC 25010", "ISO/IEC 25040", "ISO 9001"],
        "b",
        "25000 aporta vocabulario y guía general; 25040 define el proceso de evaluación; ISO 9001 "
        "es gestión de la calidad organizacional, no de producto. El modelo de características es "
        "25010.",
    ),
    (
        "Sesión 8",
        "Un usuario reporta: «Sí se puede cancelar la reserva, pero ningún compañero encuentra "
        "cómo». ¿Qué característica de calidad está afectada principalmente?",
        [
            "Adecuación funcional, subcaracterística completitud.",
            "Adecuación funcional, subcaracterística corrección.",
            "Usabilidad, subcaracterística operabilidad.",
            "Fiabilidad, subcaracterística tolerancia a fallos.",
        ],
        "c",
        "La función existe (no es completitud) y hace lo acordado (no es corrección). Lo que falla "
        "es que el usuario no logra operarla. Si la función no existiera, sería completitud; si "
        "cancelara la reserva equivocada, sería corrección.",
    ),
    (
        "Sesión 8",
        "En una prueba con 10 participantes, 7 completan la tarea. Dos abandonan (0:45 y 5:30) y "
        "uno crea la reserva para el día equivocado (3:00). ¿Cuál es la tasa de éxito y qué "
        "tiempos entran en el tiempo en tarea?",
        [
            "80 %, y entran los 8 que produjeron algún resultado.",
            "70 %, y entran solo los 7 que completaron correctamente.",
            "70 %, y entran los 10 participantes.",
            "90 %, y entran los 9 que no abandonaron.",
        ],
        "b",
        "Quien creó la reserva para el día equivocado no completó la tarea: produjo un resultado, "
        "pero no el que se pidió. La tasa es 7/10. En el tiempo en tarea entran solo quienes "
        "completaron, porque el tiempo de quien abandonó no mide eficiencia sino rendición.",
    ),
    (
        "Sesión 9",
        "Para la regla «cantidad de reservas: entero entre 1 y 5», ¿cuáles son los cuatro valores "
        "de frontera obligatorios?",
        ["1, 2, 4 y 5", "0, 1, 5 y 6", "−1, 0, 6 y 7", "1, 3 y 5"],
        "b",
        "Para toda restricción «entre A y B» las fronteras son A−1, A, B y B+1: el último inválido "
        "y el primer válido de cada extremo. Con esos cuatro se detecta cualquier error de "
        "operador de comparación (< en vez de <=).",
    ),
    (
        "Sesión 9",
        "«El archivo CSV que exporta el sistema abre correctamente en Excel y en LibreOffice, con "
        "las tildes intactas.» ¿Qué característica describe?",
        ["Portabilidad", "Compatibilidad", "Fiabilidad", "Mantenibilidad"],
        "b",
        "Hay intercambio de información con otro producto: es interoperabilidad, subcaracterística "
        "de compatibilidad. Que se mencionen dos programas no lo convierte en portabilidad; "
        "portabilidad sería que la propia aplicación se mudara a otro entorno.",
    ),
    (
        "Sesión 9",
        "Un estudiante llena un formulario de seis campos, presiona Guardar y pierde la conexión. "
        "¿Cuál es el comportamiento correcto de la aplicación?",
        [
            "Mostrar «Reserva guardada» y sincronizar cuando vuelva la señal.",
            "Mostrar «Error de red» y volver al formulario vacío.",
            "Conservar lo escrito, informar que no hay conexión y dejar el envío en estado «Pendiente» hasta confirmarlo.",
            "Reintentar en silencio con una pantalla en blanco.",
        ],
        "c",
        "Nunca se confirma una escritura que no ocurrió: si otro usuario toma el recurso primero, "
        "el falso «guardado» produce justo el conflicto que el sistema existía para evitar. La "
        "opción correcta introduce un tercer estado —pendiente— y no pierde el trabajo del usuario.",
    ),
    (
        "Sesiones 8 y 9",
        "¿Cuál de estas afirmaciones sobre la Ley de Moore es válida?",
        [
            "Garantiza que los programas se ejecutarán el doble de rápido cada dos años.",
            "Es una ley física que describe el límite de la miniaturización.",
            "Describe la duplicación aproximada del número de transistores; no la velocidad de un programa concreto.",
            "Implica que ya no hace falta optimizar el software.",
        ],
        "c",
        "Es una observación empírica de 1965 (revisada en 1975) sobre densidad de transistores. "
        "Desde ~2005 la frecuencia de reloj se estancó y los transistores adicionales van a más "
        "núcleos: un programa de un solo hilo no se acelera por sí solo. Las otras tres extienden "
        "la observación a magnitudes que no describe.",
    ),
    (
        "Sesiones 8 y 9",
        "Un estudiante entra al sistema de notas con la contraseña de un compañero, sin modificar "
        "nada. ¿Qué norma colombiana aplica principalmente?",
        [
            "Ley 1581 de 2012, porque hay datos personales involucrados.",
            "Ley 1273 de 2009, por acceso abusivo a sistema informático.",
            "Ninguna, porque no causó daño.",
            "Ley 1266 de 2008, de habeas data financiero.",
        ],
        "b",
        "El tipo penal de acceso abusivo (art. 269A) no exige que se cause daño: basta el acceso "
        "sin autorización. La 1581 obliga a quien CUSTODIA los datos; aquí el protagonista es "
        "quien ataca, así que la norma principal es la 1273.",
    ),
    (
        "Sesión 10",
        "Un algoritmo debe contar cuántas notas de un grupo son mayores o iguales a 3,0 y sumar "
        "todas las notas. ¿Qué variables necesita y cómo se inicializan?",
        [
            "Una sola variable que sirva para las dos cosas, inicializada en 0.",
            "Un contador y un acumulador, ambos inicializados en 0 antes del ciclo.",
            "Un contador inicializado en 1 y un acumulador inicializado en 0, ambos dentro del ciclo.",
            "Dos acumuladores inicializados en 0 dentro del ciclo.",
        ],
        "b",
        "El contador responde «cuántas» y crece de a 1 bajo condición; el acumulador responde "
        "«cuánto en total» y crece según el valor leído. Las dos se inicializan en 0 ANTES del "
        "ciclo: si quedan adentro se reinician en cada vuelta y el resultado siempre será el de "
        "la última nota, sin que el programa falle nunca.",
    ),
    (
        "Sesión 10",
        "En el modelo de un sistema de reservas, ¿cuál de las siguientes afirmaciones es correcta?",
        [
            "«Sala203» es una clase, porque nombra una sala concreta.",
            "«capacidad» es un método, porque describe algo de la sala.",
            "«La sala 203, con capacidad 30» es un objeto de la clase Sala, y 30 es el valor de un atributo.",
            "«duracion()» debe guardarse como atributo para no recalcularla.",
        ],
        "c",
        "Si se puede contar o señalar, es objeto; la clase es el molde. «capacidad» es atributo "
        "(sustantivo, guarda estado). Y duracion() se calcula a partir de horaInicio y horaFin: "
        "guardarla como atributo abre la posibilidad de que quede desactualizada respecto de las "
        "horas, que es la inconsistencia que la orientación a objetos existe para evitar.",
    ),
]

for i, (sesion, enunciado, opciones, correcta, justificacion) in enumerate(PREGUNTAS, start=1):
    h3(f"Pregunta {i} · {sesion}")
    p(enunciado)
    for letra, op in zip("abcd", opciones):
        marca = "  ✔" if letra == correcta else ""
        doc.add_paragraph(f"{letra}) {op}{marca}", style="List Bullet")
    par = doc.add_paragraph()
    par.add_run("Respuesta: ").bold = True
    par.add_run(f"{correcta}) ")
    par.add_run(justificacion)

# ============================== 7. Sin conexión =============================

doc.add_page_break()
h1("7. Alternativa de clase sin conexión")

p(
    "Si la sala no tiene internet o la plataforma no está disponible, el corte se puede dictar "
    "igual. Estas son las equivalencias sin computador."
)

tabla(
    ["Actividad de la aplicación", "Versión en papel o tablero"],
    [
        [
            "Ordena el desarrollo",
            "Nueve tarjetas con las acciones. Cada grupo las ordena sobre el pupitre y explica una dependencia al resto.",
        ],
        [
            "¿Quién se encarga?",
            "Dos columnas en el tablero: responsabilidades y roles. Los estudiantes pasan y trazan la línea.",
        ],
        [
            "Requisito o frase vaga",
            "Ocho frases impresas. Cada grupo las clasifica y después reescribe una vaga en el tablero.",
        ],
        [
            "Detective de calidad",
            "Ocho fallas impresas y las ocho características en el tablero. Se clasifica por votación a mano alzada.",
        ],
        [
            "Mide la usabilidad",
            "La tabla de los 10 participantes impresa. Se calcula a mano; la discusión sobre qué tiempos entran es el punto.",
        ],
        [
            "Rompe el formulario",
            "Imprimir el formulario con las reglas al lado. Cada grupo escribe tres entradas que cree que lo rompen y las intercambia con otro grupo, que actúa de validador siguiendo las reglas al pie de la letra.",
        ],
        [
            "Explora los límites",
            "Dibujar la recta numérica en el tablero y marcar entre todos los cuatro valores de frontera.",
        ],
        [
            "Simulador de duplicación",
            "La tabla lineal/exponencial en el tablero, completada a mano. El de la hoja doblada 20 veces se hace con una hoja real y una estimación a mano alzada antes de calcular.",
        ],
        [
            "¿Necesitas ese dato?",
            "Lista de nueve datos en el tablero. Cada grupo decide y, sobre todo, escribe la finalidad de los que sí pide.",
        ],
        [
            "Sigue las variables",
            "La tabla de traza dibujada en el tablero, completada fila por fila con el grupo. Funciona mejor en papel que en pantalla.",
        ],
        [
            "Arma el algoritmo / Reto integrador",
            "Pseudocódigo en el cuaderno. Se intercambian con el compañero de al lado, que ejecuta a mano los casos de prueba y anota qué salida obtiene.",
        ],
        [
            "Del objeto a la clase",
            "Tarjetas con objetos, atributos y métodos mezclados. Se agrupan sobre el pupitre.",
        ],
    ],
)

p(
    "Recomendación: cuando vuelva la conexión, pida que registren en la aplicación lo que ya "
    "resolvieron en papel. Así el panel docente refleja el avance real del grupo y el informe por "
    "concepto sigue sirviendo para el repaso de la sesión 11.",
    cursiva=True,
)

# ============================== 8. Panel docente ============================

doc.add_page_break()
h1("8. Cómo usar el panel docente")

p(
    "El panel docente está en el menú superior y solo aparece si su cuenta tiene el rol docente. "
    "Ese rol no se puede activar desde la aplicación: lo asigna la administración del curso en la "
    "base de datos, a propósito, para que ningún estudiante pueda dárselo a sí mismo."
)

h2("Qué muestra")

vinetas(
    [
        "Filtro por grupo y por estación.",
        "Lista de estudiantes con completadas, con dominio, pendientes, puntos base, nota orientativa, intentos y pistas.",
        "Detalle por estudiante: actividad por actividad, con su estado, puntaje, intentos y pistas.",
        "Actividades que más se atascan, ordenadas por porcentaje de estudiantes que la intentaron y no la completaron.",
        "Conceptos con más dificultad. Esto es una inferencia, no una medición: cuenta cuántas veces un concepto aparece etiquetado en actividades sin completar o resueltas con pistas.",
        "Exportación a CSV de lo que esté filtrado, con separador punto y coma y codificación UTF-8 con BOM, para que Excel en español lo abra directo.",
    ]
)

h2("Cómo leerlo sin sacar conclusiones de más")

vinetas(
    [
        "«Con dominio» significa resuelta al primer intento y sin pistas. Un estudiante con muchas «completadas» y pocas «dominadas» está aprendiendo, no fallando.",
        "Un número alto de intentos no es mala señal por sí solo: puede ser alguien que reintenta hasta entender. Mírelo junto con el estado final.",
        "La columna de conceptos con dificultad sirve para decidir qué repasar en clase, no para evaluar a nadie.",
        "La nota orientativa excluye desafíos opcionales, talleres y laboratorio. No es la nota del corte.",
    ]
)

h2("Rutina sugerida")

numerada(
    [
        "Antes de cada sesión: abrir el panel, filtrar por el grupo y mirar las tres actividades más atascadas de la estación anterior.",
        "Empezar la sesión resolviendo una de ellas en vivo, pidiendo el razonamiento antes de la respuesta.",
        "Antes de la sesión 11: exportar el CSV y revisar quiénes tienen menos del 40 % de avance, para acompañarlos aparte.",
    ]
)

SALIDA.parent.mkdir(parents=True, exist_ok=True)
doc.save(SALIDA)
print(f"Guía generada: {SALIDA}")
print(f"  {len(doc.paragraphs)} párrafos, {len(doc.tables)} tablas, {len(PREGUNTAS)} preguntas de parcial")
